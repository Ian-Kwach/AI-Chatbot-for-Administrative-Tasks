import json
import smtplib
import os
import shutil
import requests
import logging
import telegram
import pandas as pd
import openpyxl
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from datetime import datetime, timedelta
from googleapiclient.discovery import build
from google.oauth2 import service_account

logging.basicConfig(level=logging.INFO)

# Email Sending Function
def send_email(sender_email, sender_password, recipient, subject, body):
    try:
        server = smtplib.SMTP_SSL("smtp.gmail.com", 465)
        server.login(sender_email, sender_password)
        message = f"Subject: {subject}\n\n{body}"
        server.sendmail(sender_email, recipient, message)
        server.quit()
        return "✅ Email sent successfully!"
    except Exception as e:
        logging.error(f"❌ Error sending email: {e}")
        return "Failed to send email."

# PowerPoint Editing Function
def edit_ppt(file_path, new_text):
    if not os.path.exists(file_path):
        return f"❌ File not found: {file_path}"
    
    try:
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    shape.text = new_text  
        new_file = f"updated_{os.path.basename(file_path)}"
        prs.save(new_file)
        return f"✅ PowerPoint updated successfully: {new_file}"
    except Exception as e:
        logging.error(f"❌ Error updating PowerPoint: {e}")
        return "Failed to update PowerPoint."

# PDF Signing Function
def sign_pdf(input_pdf, output_pdf, signature_image, x=100, y=100):
    if not os.path.exists(input_pdf):
        return f"❌ File not found: {input_pdf}"

    try:
        reader = PdfReader(input_pdf)
        writer = PdfWriter()

        for page_num in range(len(reader.pages)):
            page = reader.pages[page_num]

            c = canvas.Canvas("temp_signed.pdf", pagesize=letter)
            c.drawImage(signature_image, x, y, width=200, height=50)
            c.save()

            temp_reader = PdfReader("temp_signed.pdf")
            page.merge_page(temp_reader.pages[0])

            writer.add_page(page)

        with open(output_pdf, "wb") as output:
            writer.write(output)

        return f"✅ PDF signed successfully: {output_pdf}"
    except Exception as e:
        logging.error(f"❌ Error signing PDF: {e}")
        return "Failed to sign PDF."

# File Downloading Function
def download_file(url, save_path):
    try:
        response = requests.get(url, stream=True)
        with open(save_path, "wb") as file:
            for chunk in response.iter_content(1024):
                file.write(chunk)
        return f"✅ File downloaded successfully: {save_path}"
    except Exception as e:
        logging.error(f"❌ Error downloading file: {e}")
        return "Failed to download file."

# Sending Telegram Message Function
def send_telegram_message(token, chat_id, message):
    try:
        bot = telegram.Bot(token=token)
        bot.send_message(chat_id=chat_id, text=message)
        return "✅ Telegram message sent successfully!"
    except Exception as e:
        logging.error(f"❌ Error sending Telegram message: {e}")
        return "Failed to send Telegram message."

# Scheduling Google Calendar Event
def schedule_meeting(service_account_file, calendar_id, event_title, start_time, duration=1):
    creds = service_account.Credentials.from_service_account_file(service_account_file, scopes=["https://www.googleapis.com/auth/calendar"])
    service = build("calendar", "v3", credentials=creds)

    event = {
        "summary": event_title,
        "start": {"dateTime": start_time, "timeZone": "Asia/Singapore"},
        "end": {"dateTime": (datetime.fromisoformat(start_time) + timedelta(hours=duration)).isoformat(), "timeZone": "Asia/Singapore"},
    }

    try:
        service.events().insert(calendarId=calendar_id, body=event).execute()
        return "✅ Meeting scheduled successfully!"
    except Exception as e:
        logging.error(f"❌ Error scheduling meeting: {e}")
        return "Failed to schedule meeting."

# Generating Excel Report
def generate_excel_report(data, output_file):
    try:
        df = pd.DataFrame(data)
        df.to_excel(output_file, index=False, engine="openpyxl")
        return f"✅ Excel report generated: {output_file}"
    except Exception as e:
        logging.error(f"❌ Error generating Excel report: {e}")
        return "Failed to generate Excel report."

# Chatbot Function
def chatbot():
    print("Chatbot is running! Type 'quit' to exit.")
    
    while True:
        user_input = input("You: ").strip().lower()
        if user_input == "quit":
            break

        if "send email" in user_input:
            sender_email = input("Enter your Gmail: ")
            sender_password = input("Enter your password: ")  # Use app password for security
            recipient = input("Enter recipient email: ")
            subject = input("Enter subject: ")
            body = input("Enter email body: ")
            print(send_email(sender_email, sender_password, recipient, subject, body))
            continue
        
        elif "edit ppt" in user_input:
            file_name = input("Enter PowerPoint file name (e.g., slides.pptx): ").strip()
            new_text = input("Enter new text to replace: ").strip()
            print(edit_ppt(file_name, new_text))
            continue

        elif "sign pdf" in user_input:
            file_name = input("Enter PDF file name: ").strip()
            signature = input("Enter signature image file (e.g., sign.png): ").strip()
            print(sign_pdf(file_name, f"signed_{file_name}", signature))
            continue

        elif "download file" in user_input:
            file_url = input("Enter file URL: ").strip()
            save_location = input("Enter save path (e.g., download.pdf): ").strip()
            print(download_file(file_url, save_location))
            continue

        elif "send telegram" in user_input:
            bot_token = input("Enter your Telegram Bot Token: ").strip()
            chat_id = input("Enter your Telegram Chat ID: ").strip()
            message = input("Enter message: ").strip()
            print(send_telegram_message(bot_token, chat_id, message))
            continue

        elif "schedule meeting" in user_input:
            service_account_file = input("Enter Google Calendar service account JSON file: ").strip()
            calendar_id = input("Enter Calendar ID: ").strip()
            event_title = input("Enter meeting title: ").strip()
            start_time = input("Enter start time (YYYY-MM-DDTHH:MM:SS format): ").strip()
            print(schedule_meeting(service_account_file, calendar_id, event_title, start_time))
            continue

        elif "generate report" in user_input:
            data = [
                {"Name": "John Doe", "Sales": 5000},
                {"Name": "Jane Smith", "Sales": 7000},
            ]
            output_file = input("Enter Excel report filename (e.g., report.xlsx): ").strip()
            print(generate_excel_report(data, output_file))
            continue

        else:
            print("Welcome to administrative assistant Bot!")

chatbot()
